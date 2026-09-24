# -*- coding: utf-8 -*-
import json, io, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"C:\Users\Masroor\Desktop\Venus-Aesthetics\05-content-calendar-2026\venus-calendar-app\data\calendar.json"
OUT = r"C:\Users\Masroor\Desktop\Venus-Aesthetics\05-content-calendar-2026\Venus-Content-Calendar-Sep-Dec-2026.xlsx"

D = json.load(io.open(SRC, encoding="utf-8"))
INK="000000"; GREY="F2F2F2"
SRC_FILL = {"Ready to post":"E8F3EC","Edit needed":"FBF3E6","Shoot, scripted":"EAF0F6",
            "Shoot, engaging":"F6EAF0","Shoot, b-roll":"EAF0F6","Design":"F0EDF6"}

wb = openpyxl.Workbook()
thin = Side(style="thin", color="DBDBDB"); B = Border(left=thin,right=thin,top=thin,bottom=thin)

def head(ws, freeze="C2"):
    for c in range(1, ws.max_column+1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = freeze

def widths(ws, w):
    for i,x in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width = x

def body(ws, h, bold=()):
    for i,row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=0):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border=B; cell.font=Font(size=9)
        for bi in bold:
            if bi < len(row): row[bi].font = Font(size=9, bold=True, color=INK)
        ws.row_dimensions[i+2].height = h

# ---------------------------------------------------------------- START HERE
ws = wb.active; ws.title = "START HERE"
widths(ws,[4,42,116])
s = D["stats"]
rows = [
 ("","VENUS AESTHETICS","CONTENT CALENDAR"),
 ("","Period","Monday 7 September to Sunday 6 December 2026. 13 weeks, 91 posts, one per day."),
 ("","Live board","https://venus-content-calendar-production.up.railway.app"),
 ("","",""),
 ("","WHAT IS IN EVERY POST",""),
 ("","Hook, caption, hashtags, CTA","Written to be posted as they are. Hashtags on all 91."),
 ("","A slide by slide or shot by shot spec","Carousels specify every slide: background, heading, body, footer, and a designer note. Videos specify shot order with timings, on screen text and audio."),
 ("","The state it is in","Ready to post, Edit needed, Shoot, or Design. No vague 'need making'."),
 ("","A downloadable script","Every talking video carries the full word for word script, downloadable as a text file from the board."),
 ("","A record guide","Every engaging video has a who, where, kit, beat by beat and direction guide. Most take under fifteen minutes on set."),
 ("","A format matched reference","A carousel references a real carousel. A video references a video. 29 references, all crawled with real engagement."),
 ("","",""),
 ("","THE NUMBERS",""),
 ("","Ready to post as is", str(s["readyToPost"])+" posts. Footage is finished, nothing to do."),
 ("","Edit needed", str(s["editNeeded"])+" posts. Footage exists, exact cut specified."),
 ("","To shoot", str(s["toShoot"])+" videos. These are the only things on the Production page."),
 ("","To design", str(s["toDesign"])+" carousels and statics."),
 ("","Format mix", str(s["byCreative"].get("Reel",0))+" reels, "+str(s["byCreative"].get("Carousel",0))+" carousels, "+str(s["byCreative"].get("Static",0))+" statics."),
 ("","",""),
 ("","THE RULES THIS CALENDAR FOLLOWS",""),
 ("","Nobody is quoted, anywhere","Audio in the Drive files cannot be verified, so no caption puts words in anyone's mouth. Where a patient speaks on camera the spec says subtitle their real words and cut anything unclear."),
 ("","Existing footage first", str(s["fromDrive"])+" of 91 posts are built from footage already in your Drive, each with a direct link."),
 ("","No discount posts","Zero. The only commercial posts are two free consultation statics, and the consultation is a standing service."),
 ("","Already posted content is flagged","Primelase vs Soprano, Mehwish fat freeze, Kabeer PRP, Ayesha Kamran laser, Khurasa Party Peel and Daniah Neck Lift are already live. Where reused, the spec says so and says what to change."),
 ("","Eleven new treatments","SlimFit, hyperhidrosis, masseter, lip flip, AcneOUT, Exosome Facial, CelluLITE, Scalp Detox, Fusion PRP, Derma Pen and alopecia. None had ever been scheduled."),
 ("","",""),
 ("","THE SHEETS",""),
 ("","Calendar","All 91 days with everything above."),
 ("","Production","Only the 32 videos that need shooting, numbered in the order needed."),
 ("","Scripts","The 15 full talking scripts, one per row."),
 ("","Engaging Guides","The 13 engaging videos with their record guides and references."),
 ("","Reference Library","All 29 references with format, what they demonstrate and their real engagement."),
]
for r in rows: ws.append(list(r))
ws["B1"].font = Font(bold=True, size=19, color=INK); ws["C1"].font = Font(bold=True, size=19, color="6B6B6B")
ws.row_dimensions[1].height = 28
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=3):
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        if cell.column == 2 and cell.value:
            v=str(cell.value)
            cell.font = Font(bold=True, size=12, color="6B6B6B") if v.isupper() else Font(bold=True, size=10, color=INK)

# ---------------------------------------------------------------- CALENDAR
ws = wb.create_sheet("Calendar")
ws.append(["#","Date","Day","Wk","Block","Format","Treatment","Type","STATE","OWNER",
           "HOOK","CAPTION","HASHTAGS","CTA","FOOTAGE / JOB","DRIVE LINK",
           "SPEC: slide by slide or shot by shot","SCRIPT","RECORD GUIDE",
           "REFERENCE","Ref account","Ref format","Why that reference","Ref result","Status"])
for it in D["items"]:
    ws.append([
      it["id"], it["dateLabel"], it["dayFull"], it["week"], it["block"], it["show"], it["treatment"],
      it["creative"], it["sourceLabel"], it["owner"],
      it["hook"], it["caption"], it["hashtags"], it["cta"],
      (it["driveFolder"]+" / "+it["driveFile"]) if it["fromDrive"] else it["asset"],
      it["driveUrl"] or "",
      it["spec"], it.get("script") or "", it.get("recordGuide") or "",
      it["ref"] or "", it["refAccount"] or "", it["refFormat"] or "", it["refWhy"] or "", it["refMetric"] or "", ""
    ])
widths(ws,[4,13,10,4,20,17,26,9,15,14,40,68,44,30,44,14,86,74,74,36,20,17,52,32,11])
head(ws)
for i,row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=0):
    it = D["items"][i]
    fill = SRC_FILL.get(it["sourceLabel"], "FFFFFF")
    for cell in row:
        cell.alignment = Alignment(vertical="top", wrap_text=True); cell.border=B
        cell.font = Font(size=9); cell.fill = PatternFill("solid", fgColor=fill)
    row[5].font = Font(size=9, bold=True, color=INK)
    row[8].font = Font(size=9, bold=True, color=INK)
    row[10].font = Font(size=9, bold=True)
    if it["driveUrl"]:
        row[15].hyperlink = it["driveUrl"]; row[15].value = "Open"
        row[15].font = Font(size=9, color="0000EE", underline="single", bold=True)
    if it["ref"]:
        row[19].hyperlink = it["ref"]; row[19].value = it["ref"].replace("https://","")
        row[19].font = Font(size=8, color="0000EE", underline="single")
    ws.row_dimensions[i+2].height = 150

# ---------------------------------------------------------------- PRODUCTION
ws = wb.create_sheet("Production")
ws.append(["No","Date","Wk","Kind","Video","Treatment","Hook","Shoot spec","Script","Record guide","Reference","Ref result"])
for p in D["production"]:
    ws.append([p["no"], p["date"], p["week"], p["kind"], p["title"], p["treatment"], p["hook"],
               p["spec"], p.get("script") or "", p.get("recordGuide") or "",
               p["ref"] or "", p["refMetric"] or ""])
widths(ws,[5,13,4,16,26,24,40,86,76,76,36,32]); head(ws, "C2"); body(ws, 150, bold=(4,))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[0].font = Font(size=12, bold=True, color=INK)
    if row[10].value:
        row[10].hyperlink = row[10].value
        row[10].value = str(row[10].value).replace("https://","")
        row[10].font = Font(size=8, color="0000EE", underline="single")

# ---------------------------------------------------------------- SCRIPTS
ws = wb.create_sheet("Scripts")
ws.append(["Date","Wk","Video","Treatment","Hook","FULL SCRIPT","Shoot spec"])
for it in D["items"]:
    if not it.get("script"): continue
    ws.append([it["dateLabel"], it["week"], it["asset"].replace("SHOOT: ",""), it["treatment"],
               it["hook"], it["script"], it["spec"]])
widths(ws,[13,4,24,26,40,104,80]); head(ws, "C2"); body(ws, 230, bold=(2,))

# ---------------------------------------------------------------- ENGAGING
ws = wb.create_sheet("Engaging Guides")
ws.append(["Date","Wk","Video","Hook","Caption","HOW TO RECORD IT","Spec","Reference","Ref result"])
for it in D["items"]:
    if not it.get("recordGuide"): continue
    ws.append([it["dateLabel"], it["week"], it["asset"].replace("SHOOT: ",""), it["hook"], it["caption"],
               it["recordGuide"], it["spec"], it["ref"] or "", it["refMetric"] or ""])
widths(ws,[13,4,22,40,60,100,60,36,34]); head(ws, "C2"); body(ws, 200, bold=(2,))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    if row[7].value:
        row[7].hyperlink = row[7].value
        row[7].value = str(row[7].value).replace("https://","")
        row[7].font = Font(size=8, color="0000EE", underline="single")

# ---------------------------------------------------------------- REFERENCES
ws = wb.create_sheet("Reference Library")
ws.append(["Key","Account","Format","What it demonstrates","Real result","Used on","Link"])
for r in D["referenceLibrary"]:
    ws.append([r["key"].replace("_"," "), r["account"], r["format"], r["why"], r["metric"], r["usedOn"], r["url"]])
widths(ws,[24,22,20,78,44,9,42]); head(ws, "C2"); body(ws, 52, bold=(0,))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    if row[6].value:
        row[6].hyperlink = row[6].value
        row[6].value = str(row[6].value).replace("https://","")
        row[6].font = Font(size=9, color="0000EE", underline="single")

wb.save(OUT)
print("saved:", OUT)
print("sheets:", wb.sheetnames)

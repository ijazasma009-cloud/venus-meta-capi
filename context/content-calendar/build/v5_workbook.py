# -*- coding: utf-8 -*-
import json, io, os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = r"C:\Users\Masroor\Desktop\Venus-Aesthetics\05-content-calendar-2026"
SRC = os.path.join(ROOT, "venus-calendar-app", "data", "calendar.json")
OUT = os.path.join(ROOT, "Venus-Content-Calendar-8-Weeks.xlsx")

D = json.load(io.open(SRC, encoding="utf-8")); s = D["stats"]
INK = "000000"
FILL = {"Ready to post": "E8F3EC", "Edit needed": "FBF3E6", "Shoot, props": "EAF0F6",
        "Shoot, engaging": "F6EAF0", "Design": "F0EDF6"}

wb = openpyxl.Workbook()
thin = Side(style="thin", color="DBDBDB"); B = Border(left=thin, right=thin, top=thin, bottom=thin)

def head(ws, freeze="C2"):
    for c in range(1, ws.max_column + 1):
        k = ws.cell(row=1, column=c)
        k.fill = PatternFill("solid", fgColor=INK); k.font = Font(bold=True, color="FFFFFF", size=10)
        k.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 34; ws.freeze_panes = freeze

def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x

def body(ws, h, bold=()):
    for n, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row)):
        for k in row:
            k.alignment = Alignment(vertical="top", wrap_text=True); k.border = B; k.font = Font(size=9)
        for bi in bold:
            if bi < len(row): row[bi].font = Font(size=9, bold=True, color=INK)
        ws.row_dimensions[n + 2].height = h

# ------------------------------------------------------------------ START HERE
ws = wb.active; ws.title = "START HERE"
widths(ws, [4, 42, 120])
rows = [
 ("", "VENUS AESTHETICS", "8 WEEK CONTENT CALENDAR"),
 ("", "Period", D["period"]["label"] + ".  %d posts, one a day." % s["total"]),
 ("", "Live board", "https://venus-content-calendar-production.up.railway.app"),
 ("", "", ""),
 ("", "HOW THIS ONE IS DIFFERENT", ""),
 ("", "Every post is an idea, not a category",
  "Each row carries a named creative concept and the caption is built around it. A receipt for a year of "
  "shaving. A jar of retired razors at reception. A live clock running through a treatment. If a caption "
  "would read the same with the idea removed, it was rewritten."),
 ("", "Three voices, used deliberately",
  "Warm on treatment and conversion posts. Useful on educational posts. Playful on the engaging ones. "
  "Same brand, different energy depending on the job the post is doing."),
 ("", "No statistics in captions",
  "Research sits behind the calendar, never in front of the customer. No AQI readings, no study "
  "citations, no regulator figures in a caption. Those inform the idea and then get out of the way."),
 ("", "No Venus Glow and no facials", "Removed entirely at the client's instruction. Five clips were pulled from the pool."),
 ("", "", ""),
 ("", "THE WEEK", ""),
 ("", "Monday, Treatment Film", "A treatment you already filmed, shot through a different device each week. A timer, a split screen, the client's eyeline, the machine's eye view."),
 ("", "Tuesday, Skin School", "Education that does not look like education. A receipt, a bus timetable, a nutrition label, a game select screen."),
 ("", "Wednesday, Ask Venus", "Eight unused Dr. Uzair FAQs, each opening on a different device so they never feel like a row of talking heads."),
 ("", "Thursday, The Results", "Proof without needing skin you have no consent for. A jar of razors, a tape measure, a hairbrush, a kameez that fits differently."),
 ("", "Friday, Real Talk", "A quick shoot, under 25 minutes, phone only. Rating hair removal methods, myth or truth, guess the treatment from the sound."),
 ("", "Saturday, The Team", "Clips you already have, re-cut. Three carry spelling fixes that have been sitting in the footage."),
 ("", "Sunday, Your Glow Plan", "Built to be screenshotted and forwarded. A calendar annotated in pen, a departure board, a metro map of your sessions."),
 ("", "", ""),
 ("", "THE NUMBERS", ""),
 ("", "Ready to post", "%d posts. The Dr. Uzair FAQs, already subtitled in English." % s["readyToPost"]),
 ("", "Edit needed", "%d posts. Footage exists, the edit is written out step by step." % s["editNeeded"]),
 ("", "To shoot", "%d. All short, all with a prop list and a beat by beat guide." % s["toShoot"]),
 ("", "To design", "%d graphics and carousels, each written slide by slide." % s["toDesign"]),
 ("", "Format mix", "%d reels, %d carousels and graphics. No single images."
   % (s["byCreative"].get("Reel", 0), s["byCreative"].get("Carousel", 0))),
 ("", "Mix by type", "Treatment %d, Educational %d, Engaging %d, Conversion %d."
   % (s["byBucket"].get("Treatment", 0), s["byBucket"].get("Educational", 0),
      s["byBucket"].get("Engaging", 0), s["byBucket"].get("Conversion", 0))),
 ("", "Your own footage", "%d of %d posts point at a real file in your Drive." % (s["driveLinked"], s["total"])),
 ("", "", ""),
 ("", "THE SHEETS", ""),
 ("", "Calendar", "Every post with its idea, caption, hashtags, the full brief and the reference."),
 ("", "Production", "Only the shoots. Numbered, with props and a beat by beat guide."),
 ("", "Designer Briefs", "Every designed post, slide by slide, in numbered pointers."),
 ("", "Reference Library", "Every reference with its real engagement. A reel references a reel."),
]
for r in rows: ws.append(list(r))
ws["B1"].font = Font(bold=True, size=19, color=INK); ws["C1"].font = Font(bold=True, size=19, color="6B6B6B")
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=3):
    for k in row:
        k.alignment = Alignment(vertical="top", wrap_text=True)
        if k.column == 2 and k.value:
            v = str(k.value)
            k.font = Font(bold=True, size=12, color="6B6B6B") if v.isupper() else Font(bold=True, size=10, color=INK)

# ------------------------------------------------------------------- CALENDAR
ws = wb.create_sheet("Calendar")
ws.append(["#", "Date", "Day", "Wk", "TYPE", "FORMAT", "THE IDEA", "Concept", "Treatment", "Post",
           "Voice", "STATE", "Owner", "HOOK, on screen", "CAPTION", "HASHTAGS",
           "FOOTAGE / JOB", "Secs", "LINK", "THE BRIEF, numbered", "SHOOT GUIDE",
           "REFERENCE", "Ref account", "What to copy", "Ref result", "Status"])
for i in D["items"]:
    ws.append([i["id"], i["dateLabel"], i["dayFull"], i["week"], i["bucket"], i["show"],
               i["conceptLine"], i["variant"], i["treatment"], i["creative"], i["voice"],
               i["sourceLabel"], i["owner"], i["hook"], i["caption"], i["hashtags"],
               (i["driveFolder"] + " / " + i["driveFile"]) if i["fromDrive"] else i["asset"],
               i.get("sourceSeconds") or "", i["driveUrl"] or "", i["spec"], i["recordGuide"],
               i["ref"] or "", i["refAccount"] or "", i["refWhy"] or "", i["refMetric"] or "", ""])
widths(ws, [4, 13, 10, 4, 13, 16, 52, 26, 22, 9, 9, 14, 11, 40, 66, 34, 38, 7, 8, 94, 76, 30, 16, 46, 30, 10])
head(ws)
for n, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row)):
    i = D["items"][n]; f = FILL.get(i["sourceLabel"], "FFFFFF")
    for k in row:
        k.alignment = Alignment(vertical="top", wrap_text=True); k.border = B
        k.font = Font(size=9); k.fill = PatternFill("solid", fgColor=f)
    for c in (4, 5, 11): row[c].font = Font(size=9, bold=True, color=INK)
    row[6].font = Font(size=9, bold=True, italic=True, color=INK)
    row[13].font = Font(size=9, bold=True)
    if i["driveUrl"]:
        row[18].hyperlink = i["driveUrl"]; row[18].value = "Open"
        row[18].font = Font(size=9, color="0000EE", underline="single", bold=True)
    if i["ref"]:
        row[21].hyperlink = i["ref"]; row[21].value = i["ref"].replace("https://", "")
        row[21].font = Font(size=8, color="0000EE", underline="single")
    ws.row_dimensions[n + 2].height = 170

# ----------------------------------------------------------------- PRODUCTION
ws = wb.create_sheet("Production")
ws.append(["No", "Date", "Wk", "Kind", "Concept", "THE IDEA", "Hook", "HOW TO SHOOT IT", "The brief", "Reference", "Ref result"])
for p in D["production"]:
    ws.append([p["no"], p["date"], p["week"], p["kind"], p["variant"], p["conceptLine"], p["hook"],
               p["recordGuide"], p["spec"], p["ref"] or "", p["refMetric"] or ""])
widths(ws, [5, 13, 4, 16, 28, 54, 40, 84, 84, 30, 30]); head(ws); body(ws, 180, bold=(4,))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[0].font = Font(size=13, bold=True, color=INK)
    row[5].font = Font(size=9, bold=True, italic=True)
    if row[9].value:
        row[9].hyperlink = row[9].value; row[9].value = str(row[9].value).replace("https://", "")
        row[9].font = Font(size=8, color="0000EE", underline="single")

# ------------------------------------------------------------ DESIGNER BRIEFS
ws = wb.create_sheet("Designer Briefs")
ws.append(["Date", "Wk", "Concept", "THE IDEA", "Hook", "SLIDE BY SLIDE", "Pull frames from", "Reference"])
for i in D["items"]:
    if i["source"] != "DESIGN": continue
    ws.append([i["dateLabel"], i["week"], i["variant"], i["conceptLine"], i["hook"],
               i["spec"], i.get("frameSource", ""), i["ref"] or ""])
widths(ws, [13, 4, 28, 54, 40, 104, 32, 30]); head(ws); body(ws, 230, bold=(2,))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    row[3].font = Font(size=9, bold=True, italic=True)
    if row[7].value:
        row[7].hyperlink = row[7].value; row[7].value = str(row[7].value).replace("https://", "")
        row[7].font = Font(size=8, color="0000EE", underline="single")

# ----------------------------------------------------------------- REFERENCES
ws = wb.create_sheet("Reference Library")
ws.append(["Account", "Format", "What to copy from it", "Real result", "Used on", "Link"])
for r in D["referenceLibrary"]:
    ws.append([r["account"], r["format"], r["why"], r["metric"], r["usedOn"], r["url"]])
widths(ws, [22, 14, 82, 46, 9, 42]); head(ws); body(ws, 56, bold=(0,))
for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
    if row[5].value:
        row[5].hyperlink = row[5].value; row[5].value = str(row[5].value).replace("https://", "")
        row[5].font = Font(size=9, color="0000EE", underline="single")

wb.save(OUT)
print("saved:", OUT); print("sheets:", wb.sheetnames)
